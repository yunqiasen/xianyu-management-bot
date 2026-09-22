"""S1: typed defaults are materialized into the same account configuration consumers read."""
import json
from fastapi import FastAPI
import httpx
from sqlalchemy import select
from common.models.user import User, UserRole
from common.models.admin_control import AdminAudit
from common.models.system_setting import SystemSetting
from common.models.user_setting import UserSetting
from common.models.xy_account import XYAccount
from common.models.xy_delivery_block_rule import XYDeliveryBlockRule
from common.services import account_policy
from test_runtime import DatabaseCase


class ConfigInheritanceTests(DatabaseCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        from app.api import deps
        from app.api.routes import cookies, user_settings, system_settings
        from app.services.account_service import AccountService
        from app.services.system_setting_service import SystemSettingService
        async with self.engine.begin() as connection:
            await connection.run_sync(SystemSetting.__table__.create)
            await connection.run_sync(UserSetting.__table__.create)
            await connection.run_sync(AdminAudit.__table__.create)
            await connection.run_sync(XYDeliveryBlockRule.__table__.create)
        async with self.factory() as db:
            db.add(SystemSetting(key='xymb.account.defaults',value='{"version":0,"values":{}}'))
            db.add(UserSetting(user_id=7,key='legacy.unknown',value='false'))
            await db.commit()
        self.owner=7;self.role=UserRole.MEMBER
        async def user(): return User(id=self.owner,username='fixture',role=self.role)
        async def service():
            async with self.factory() as db: yield AccountService(db)
        async def db_session():
            async with self.factory() as db: yield db
        async def settings_service():
            async with self.factory() as db: yield SystemSettingService(db)
        self.app=FastAPI();self.app.include_router(cookies.router,prefix='/cookies')
        self.app.include_router(user_settings.router,prefix='/user-settings')
        self.app.include_router(system_settings.router,prefix='/system-settings')
        self.app.dependency_overrides[deps.get_current_active_user]=user
        self.app.dependency_overrides[deps.get_account_service]=service
        self.app.dependency_overrides[deps.get_db_session]=db_session
        self.app.dependency_overrides[deps.get_system_setting_service]=settings_service
        self.client=httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app),base_url='http://fixture')
        self.base='/cookies/fixture/configuration'

    async def asyncTearDown(self):
        await self.client.aclose();await super().asyncTearDown()

    async def read(self):
        response=await self.client.get(self.base+'/effective')
        self.assertEqual(response.status_code,200,response.text)
        return response.json()['data']

    async def change(self,scope,key,value=None,action='set',version=None):
        state=await self.read()
        response=await self.client.put(self.base+'/settings',json={'scope':scope,'key':key,'action':action,
            'value':value,'expected_version':state['versions'][scope] if version is None else version})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()['data']

    async def test_zero_false_empty_null_and_inherit_are_distinct(self):
        state=await self.change('account','pause_duration',0)
        self.assertEqual(state['values']['pause_duration'],{'value':0,'source':'account','inherited':False})
        await self.change('user','pause_duration',15)
        self.assertEqual((await self.read())['values']['pause_duration']['value'],0)
        state=await self.change('account','pause_duration',action='inherit')
        self.assertEqual(state['values']['pause_duration']['value'],15)
        self.assertEqual(state['values']['pause_duration']['source'],'user')
        self.role=UserRole.ADMIN
        await self.change('system','pause_duration',30)
        state=await self.change('user','pause_duration',action='inherit')
        self.assertEqual(state['values']['pause_duration']['value'],30)
        self.assertEqual(state['values']['pause_duration']['source'],'system')
        await self.change('user','auto_polish',True)
        state=await self.change('account','auto_polish',False)
        self.assertIs(state['values']['auto_polish']['value'],False)
        state=await self.change('account','delivery_disabled_reason','')
        self.assertEqual(state['values']['delivery_disabled_reason']['value'],'')
        state=await self.change('account','risk',None)
        self.assertIsNone(state['values']['risk']['value'])
        async with self.factory() as db:
            account=await db.get(XYAccount,1)
            self.assertEqual(account.pause_duration,30)
            self.assertFalse(account.auto_polish)
            self.assertEqual(account.cookie,'unb=101; token=old')
            self.assertEqual(await db.scalar(select(UserSetting.value).where(UserSetting.key=='legacy.unknown')),'false')

    async def test_risk_inheritance_changes_the_runtime_budget_and_keeps_manual_disable(self):
        from common.services.account_request_budget import account_risk_config,RequestBudgetPolicy
        async with self.factory() as db:
            account=await db.get(XYAccount,1);account.status='disabled';account.disable_reason='manual'
            await db.commit()
        await self.change('user','risk',{'min_interval_seconds':17})
        state=await self.change('account','risk',action='inherit')
        self.assertEqual(state['values']['risk']['source'],'user')
        async with self.factory() as db:
            account=await db.get(XYAccount,1)
            self.assertEqual(RequestBudgetPolicy.from_risk_config(account_risk_config(account)).interval_ms,17000)
            self.assertEqual((account.status,account.disable_reason),('disabled','manual'))
            self.assertCountEqual(account_policy.pending_consumers(account),['scheduler','websocket'])

    async def test_legacy_field_update_becomes_explicit_and_fences_old_configuration(self):
        await self.change('user','pause_duration',18)
        before=await self.change('account','pause_duration',action='inherit')
        response=await self.client.put('/cookies/fixture/pause-duration',json={'pause_duration':0})
        self.assertEqual(response.status_code,200,response.text)
        after=await self.read()
        self.assertGreater(after['versions']['account'],before['versions']['account'])
        self.assertEqual(after['values']['pause_duration']['source'],'account')
        self.assertEqual(after['values']['pause_duration']['value'],0)
        await self.change('user','pause_duration',25)
        self.assertEqual((await self.read())['values']['pause_duration']['value'],0)

    async def test_scope_conflicts_and_types_reject_without_writes(self):
        initial=await self.read()
        for key,value in (('auto_polish','false'),('pause_duration',False),('pause_duration',-1),('unknown',0)):
            response=await self.client.put(self.base+'/settings',json={'scope':'account','key':key,'action':'set','value':value,'expected_version':0})
            self.assertIn(response.status_code,(400,422),response.text)
        self.assertEqual((await self.read())['versions'],initial['versions'])
        response=await self.client.put(self.base+'/settings',json={'scope':'system','key':'pause_duration','action':'set','value':5,'expected_version':0})
        self.assertEqual(response.status_code,403)
        await self.change('account','pause_duration',3)
        response=await self.client.put(self.base+'/settings',json={'scope':'account','key':'pause_duration','action':'set','value':4,'expected_version':0})
        self.assertEqual(response.status_code,409)
        self.owner=8
        self.assertEqual((await self.client.get(self.base+'/effective')).status_code,404)

    async def test_legacy_endpoints_preserve_dedicated_defaults(self):
        self.role=UserRole.ADMIN
        await self.change('system','pause_duration',20)
        await self.change('user','pause_duration',12)
        before=await self.read()
        for method,path in (('put','/system-settings/xymb.account.defaults'),
                            ('put','/user-settings/xymb.account.defaults'),
                            ('delete','/user-settings/xymb.account.defaults')):
            response=await self.client.request(method,path,**({'json':{'value':'{}'}} if method=='put' else {}))
            self.assertEqual(response.status_code,409,response.text)
        self.assertEqual(await self.read(),before)
        response=await self.client.put('/user-settings/legacy.unknown',json={'value':'0'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual((await self.client.get('/user-settings/legacy.unknown')).json()['value'],'0')

    async def test_legacy_login_and_delivery_edits_update_the_shared_configuration_version(self):
        from common.models.xy_delivery_block_rule import XYDeliveryBlockRule
        async with self.engine.begin() as connection:
            await connection.run_sync(XYDeliveryBlockRule.__table__.create, checkfirst=True)
        before=await self.change('account','show_browser',action='inherit')
        result=await self.client.put('/cookies/fixture/login-info',json={'show_browser':False})
        self.assertEqual(result.status_code,200,result.text)
        state=await self.read()
        self.assertGreater(state['versions']['account'],before['versions']['account'])
        self.assertEqual(state['values']['show_browser']['source'],'account')
        await self.change('user','delivery_disabled_reason','default reason')
        before=await self.change('account','delivery_disabled_reason',action='inherit')
        result=await self.client.put('/cookies/fixture/delivery-disabled',json={
            'delivery_disabled':True,'delivery_disabled_reason':'explicit reason'})
        self.assertEqual(result.status_code,200,result.text)
        state=await self.read()
        self.assertGreater(state['versions']['account'],before['versions']['account'])
        self.assertEqual(state['values']['delivery_disabled_reason']['value'],'explicit reason')
        self.assertEqual(state['values']['delivery_disabled_reason']['source'],'account')
        result=await self.client.put('/cookies/fixture/delivery-block-rules',json={
            'rules':[{'rule_code':'buyer_credit_zero','enabled':False}]})
        self.assertEqual(result.status_code,200,result.text)
        self.assertGreater((await self.read())['versions']['account'],state['versions']['account'])

    async def test_existing_pause_range_is_supported_in_typed_and_legacy_forms(self):
        await self.change('account','pause_duration',1800)
        response=await self.client.put('/cookies/fixture/pause-duration',json={'pause_duration':3600})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual((await self.read())['values']['pause_duration']['value'],3600)

    async def test_excel_switches_are_explicit_and_missing_columns_keep_their_source(self):
        from io import BytesIO
        from openpyxl import Workbook
        from unittest.mock import AsyncMock, patch
        from common.models.card import Card
        await self.create_owner()
        async with self.engine.begin() as connection:
            await connection.run_sync(Card.__table__.create)
        await self.change('user','auto_polish',True)
        await self.change('account','auto_polish',action='inherit')
        await self.change('user','auto_red_flower',True)
        await self.change('account','auto_red_flower',action='inherit')
        book=Workbook();sheet=book.active;sheet.title='账号基本信息'
        sheet.append(['账号ID','Cookie']);sheet.append(['fixture','********'])
        sheet=book.create_sheet('账号开关配置');sheet.append(['账号ID','商品擦亮']);sheet.append(['fixture','false'])
        data=BytesIO();book.save(data);book.close()
        with patch('app.services.account_jobs.async_session_maker',self.factory), \
             patch('app.services.account_jobs.validate_credentials',AsyncMock(return_value='unb=101; token=new')), \
             patch('app.services.account_jobs.restart_account',AsyncMock()):
            response=await self.client.post('/cookies/import',files={'file':('settings.xlsx',data.getvalue())})
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(response.json()['success'],response.text)
        job=response.json()['data']['credential_jobs'][0]
        self.assertEqual((await self.client.get('/cookies/fixture/credential-jobs/'+job['id'])).json()['data']['status'],'verified')
        state=await self.read()
        self.assertEqual(state['values']['auto_polish'],{'value':False,'source':'account','inherited':False})
        self.assertEqual(state['values']['auto_red_flower'],{'value':True,'source':'user','inherited':True})

    async def test_typed_delivery_reason_is_projected_into_the_active_rule(self):
        response=await self.client.put('/cookies/fixture/delivery-disabled',json={
            'delivery_disabled':True,'delivery_disabled_reason':'before'})
        self.assertEqual(response.status_code,200,response.text)
        await self.change('user','delivery_disabled_reason','after')
        await self.change('account','delivery_disabled_reason',action='inherit')
        response=await self.client.get('/cookies/fixture/delivery-block-rules')
        rules=response.json()['data']
        self.assertEqual(next(row for row in rules if row['rule_code']=='buyer_credit_zero')['block_reason'],'after')

    async def import_workbook(self, basic_headers, basic_values, switches=None):
        from io import BytesIO
        from openpyxl import Workbook
        from unittest.mock import AsyncMock, patch
        from common.models.card import Card
        await self.create_owner()
        async with self.engine.begin() as connection:
            await connection.run_sync(Card.__table__.create, checkfirst=True)
        book = Workbook(); sheet = book.active; sheet.title = '账号基本信息'
        sheet.append(['账号ID', 'Cookie', *basic_headers]); sheet.append(['fixture', '********', *basic_values])
        if switches:
            sheet = book.create_sheet('账号开关配置')
            sheet.append(['账号ID', *switches.keys()]); sheet.append(['fixture', *switches.values()])
        data = BytesIO(); book.save(data); book.close()
        with patch('app.services.account_jobs.async_session_maker', self.factory), \
             patch('app.services.account_jobs.validate_credentials', AsyncMock(return_value='unb=101; token=imported')), \
             patch('app.services.account_jobs.restart_account', AsyncMock()):
            return await self.client.post('/cookies/import', files={'file': ('settings.xlsx', data.getvalue())})

    async def test_excel_legacy_delivery_switch_updates_the_actual_rule(self):
        response = await self.client.put('/cookies/fixture/delivery-disabled', json={
            'delivery_disabled': True, 'delivery_disabled_reason': 'old rule', 'auto_close_order': True,
            'delivery_only_card_after_close': True, 'delivery_disabled_excluded_items': ['item-1']})
        self.assertEqual(response.status_code, 200, response.text)
        response = await self.import_workbook([], [], {'禁止发货': 'false'})
        self.assertTrue(response.json()['success'], response.text)
        rules = (await self.client.get('/cookies/fixture/delivery-block-rules')).json()['data']
        rule = next(row for row in rules if row['rule_code'] == 'buyer_credit_zero')
        self.assertFalse(rule['enabled'], rule)
        self.assertFalse(rule['auto_close_order'], rule)
        self.assertFalse(rule['only_card_after_close'], rule)
        self.assertEqual(rule['block_reason'], '', rule)  # legacy rule API uses an empty display string
        self.assertIsNone((await self.read())['values']['delivery_disabled_reason']['value'])

    async def test_excel_profile_rejects_out_of_range_without_replacing_credentials(self):
        before = await self.read()
        response = await self.import_workbook(['暂停时长(秒)'], [3601])
        self.assertFalse(response.json()['success'], response.text)
        self.assertEqual(response.json()['data']['submitted'], 0)
        self.assertEqual((await self.read())['values']['pause_duration'], before['values']['pause_duration'])
        async with self.factory() as db:
            self.assertEqual((await db.get(XYAccount, 1)).cookie, 'unb=101; token=old')
