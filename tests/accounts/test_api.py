"""S1：ASGI HTTP→真实路由/服务→临时数据库；平台验证为 S2 接缝。"""
from pathlib import Path
import sys
import unittest
from unittest.mock import AsyncMock, patch
sys.path.insert(0, str(Path(__file__).parents[2] / 'backend-web'))
from fastapi import FastAPI
import httpx
from common.models.xy_account import XYAccount
from common.models.user import User
from test_runtime import DatabaseCase


class AccountApiTests(DatabaseCase):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        from app.core.config import get_settings
        self.token_setting = patch.object(get_settings(), 'internal_api_token', 'fixture-' * 8)
        self.token_setting.start()
        self.addCleanup(self.token_setting.stop)
        from app.api import deps
        from app.api.routes import cookies
        from app.services.account_service import AccountService
        self.app=FastAPI(); self.app.include_router(cookies.router, prefix='/cookies')
        self.owner=7
        async def user(): return User(id=self.owner, username='fixture')
        async def db():
            async with self.factory() as session: yield session
        async def svc():
            async with self.factory() as session: yield AccountService(session)
        self.app.dependency_overrides[deps.get_current_active_user]=user
        self.app.dependency_overrides[deps.get_account_service]=svc
        self.app.dependency_overrides[deps.get_db_session]=db
        self.client=httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app),base_url='http://fixture')

    async def asyncTearDown(self):
        await self.client.aclose(); await super().asyncTearDown()

    async def test_runtime_read_ownership_and_secret_redaction(self):
        response=await self.client.get('/cookies/fixture/runtime')
        self.assertEqual(response.status_code,200)
        self.assertNotIn('secret', response.text)
        self.assertNotIn('token=old', response.text)
        self.owner=8
        self.assertEqual((await self.client.get('/cookies/fixture/runtime')).status_code,404)

    async def test_job_start_verified_and_profile_preserved(self):
        with (patch('app.services.account_jobs.validate_credentials', AsyncMock(return_value='unb=101; token=new')),
             patch('app.services.account_jobs.async_session_maker', self.factory),
             patch('app.services.account_jobs.restart_account', AsyncMock())):
            response=await self.client.post('/cookies/fixture/credential-jobs',json={'value':'unb=101; token=new','expected_version':0})
        self.assertEqual(response.status_code,202, response.text)
        job=response.json()['data']
        result=await self.client.get(f"/cookies/fixture/credential-jobs/{job['id']}")
        self.assertEqual(result.json()['data']['status'],'verified')
        async with self.factory() as db:
            a=await db.get(XYAccount,1)
            self.assertEqual(a.login_password,'secret')
            self.assertEqual(a.cookie,'unb=101; token=new')

    async def test_old_job_poll_and_cancel_preserve_the_new_session(self):
        from common.services import account_policy as policy
        for version in ('credential_version', 'config_version', 'generation'):
            for now in (200, 401):
                for action in ('get', 'delete', 'runtime'):
                    with self.subTest(version=version, now=now, action=action):
                        async with self.factory() as session:
                            account = await session.get(XYAccount, 1)
                            account.metadata_json = {}
                            job = policy.start_job(account, owner_id=7, kind='cookie_import', now=100)
                            state = policy.snapshot(account)
                            state[version] += 1
                            state.update(business_state='ready', reason=None, last_success_at=150)
                            policy.store(account, state)
                            await session.commit()
                        with patch('common.services.account_policy.time.time', return_value=now):
                            url = '/cookies/fixture/runtime' if action == 'runtime' else (
                                '/cookies/fixture/credential-jobs/' + job['id'])
                            response = await self.client.request('DELETE' if action == 'delete' else 'GET', url)
                            self.assertEqual(response.status_code, 200, response.text)
                            runtime = (await self.client.get('/cookies/fixture/runtime')).json()['data']
                        self.assertEqual(runtime['business_state'], 'ready')
                        self.assertIsNone(runtime['reason'])
                        self.assertEqual(runtime['last_success_at'], 150)
                        self.assertEqual(runtime[version], 1)
                        self.assertEqual(runtime['jobs'][0]['status'], 'superseded')

    async def test_new_import_does_not_reuse_a_job_from_an_old_binding(self):
        from common.services import account_policy as policy
        for version in ('credential_version', 'config_version', 'generation'):
            with self.subTest(version=version):
                async with self.factory() as session:
                    account = await session.get(XYAccount, 1)
                    account.metadata_json = {}
                    old = policy.start_job(account, owner_id=7, kind='verification', now=100)
                    state = policy.snapshot(account)
                    state[version] += 1
                    state.update(business_state='ready', reason=None)
                    policy.store(account, state)
                    await session.commit()
                with (patch('common.services.account_policy.time.time', return_value=200),
                      patch('app.services.account_jobs.validate_credentials', AsyncMock(return_value='unb=101; token=new')),
                      patch('app.services.account_jobs.async_session_maker', self.factory),
                      patch('app.services.account_jobs.restart_account', AsyncMock())):
                    response = await self.client.post('/cookies/fixture/credential-jobs', json={
                        'value': 'unb=101; token=new', 'expected_version': state['credential_version']})
                    self.assertEqual(response.status_code, 202, response.text)
                    job = response.json()['data']
                    self.assertNotEqual(job['id'], old['id'])
                    self.assertEqual((await self.client.get(
                        '/cookies/fixture/credential-jobs/' + job['id'])).json()['data']['status'], 'verified')
                    self.assertEqual((await self.client.get(
                        '/cookies/fixture/credential-jobs/' + old['id'])).json()['data']['status'], 'superseded')

    async def test_expired_job_poll_does_not_pause_its_replacement(self):
        from app.services.account_service import AccountService
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            service = AccountService(session)
            with patch('common.services.account_policy.time.time', return_value=100):
                old = await service.start_credential_job(account, 'cookie_import', 7)
            with patch('common.services.account_policy.time.time', return_value=401):
                new = await service.start_credential_job(account, 'cookie_import', 7)
        self.assertNotEqual(old['id'], new['id'])
        with patch('common.services.account_policy.time.time', return_value=402):
            runtime = (await self.client.get('/cookies/fixture/runtime')).json()['data']
        self.assertEqual(runtime['business_state'], 'recovering')
        self.assertEqual({job['id']: job['status'] for job in runtime['jobs']},
                         {old['id']: 'expired', new['id']: 'processing'})

    async def test_qr_status_and_cookie_are_owner_scoped(self):
        from app.api.routes import qr_login
        self.app.include_router(qr_login.router)
        qr_login.SESSION_OWNER['fixture-session']=8
        for path in ('status','cookie'):
            response=await self.client.get(f'/qr-login/{path}/fixture-session')
            self.assertEqual(response.status_code,404)
        qr_login.SESSION_OWNER.pop('fixture-session',None)

    async def test_details_never_return_login_password(self):
        # schema serialization alone must redact even legacy builders passing secrets.
        from common.schemas.account import AccountDetail
        payload=AccountDetail(pk=1,id='fixture',value='unb=101; token=old',enabled=True,
                              auto_confirm=False,login_password='secret')
        self.assertNotIn('secret',payload.model_dump_json())
        self.assertNotIn('token=old',payload.model_dump_json())

    async def test_qr_refresh_binds_proxy_and_preserves_disabled_profile(self):
        from app.api.routes import qr_login
        self.app.include_router(qr_login.router)
        async with self.factory() as db:
            a=await db.get(XYAccount,1); a.status='disabled'
            a.proxy_type='http'; a.proxy_host='127.0.0.1'; a.proxy_port=8123
            await db.commit()
        generator=AsyncMock(return_value={'success':True,'session_id':'qr-fixture','qr_code_url':'fixture'})
        with patch.object(qr_login.qr_login_manager,'generate_qr_code',generator):
            response=await self.client.post('/qr-login/generate',json={'account_id':'fixture'})
        self.assertEqual(response.status_code,200,response.text)
        generator.assert_awaited_once_with(proxy='http://127.0.0.1:8123')
        async with self.factory() as db:
            a=await db.get(XYAccount,1)
            self.assertEqual(a.status,'disabled')
            self.assertEqual(a.login_password,'secret')

    async def test_proxy_secret_preserved_and_version_incremented(self):
        from app.api.routes import proxy
        self.app.include_router(proxy.router,prefix='/proxy')
        async with self.factory() as db:
            a=await db.get(XYAccount,1);a.proxy_pass='secret';await db.commit()
        response=await self.client.put('/proxy/fixture',json={'proxy_type':'http','proxy_host':'127.0.0.1',
               'proxy_port':8123,'proxy_user':'u','proxy_pass':'********'})
        self.assertTrue(response.json()['success'])
        response=await self.client.get('/proxy/fixture')
        self.assertNotIn('secret',response.text)
        async with self.factory() as db:
            from common.services import account_policy as p
            a=await db.get(XYAccount,1)
            self.assertEqual(a.proxy_pass,'secret')
            self.assertEqual(p.snapshot(a)['config_version'],1)

    async def test_legacy_cookie_update_masks_are_noop(self):
        response=await self.client.put('/cookies/fixture',json={'value':'********'})
        self.assertEqual(response.status_code,200,response.text)
        async with self.factory() as db:
            self.assertEqual((await db.get(XYAccount,1)).cookie,'unb=101; token=old')

    async def test_unknown_browser_session_is_not_proxied(self):
        from app.api.routes import password_login
        self.app.include_router(password_login.router)
        # 未知会话必须在外部服务调用前结束。
        with patch.object(password_login.httpx, 'AsyncClient') as remote:
            remote.return_value.__aenter__.return_value.get=AsyncMock(side_effect=RuntimeError('fixture'))
            result=await self.client.get('/password-login/check/not-owned-fixture')
        self.assertEqual(result.status_code,404)

    async def test_browser_control_api_scoped_and_no_store_screenshot(self):
        from app.api.routes import password_login as route
        from common.services.account_browser import AccountBrowserSessions
        self.app.include_router(route.router)
        manager=AccountBrowserSessions(self.factory)
        context=AsyncMock();page=AsyncMock();page.mouse=AsyncMock();page.screenshot.return_value=b'png'
        job=await manager.bind('fixture',7,context,page)
        with patch.object(route,'verification_sessions',manager,create=True):
            shot=await self.client.get(f"/password-login/verification/{job['id']}/screenshot")
            self.assertEqual(shot.status_code,200)
            self.assertEqual(shot.headers['cache-control'],'no-store')
            self.owner=8
            self.assertEqual((await self.client.post(f"/password-login/verification/{job['id']}/control",json={'action':'click','x':1,'y':2})).status_code,404)
            self.owner=7
            self.assertEqual((await self.client.post(f"/password-login/verification/{job['id']}/control",json={'action':'click','x':1,'y':2})).status_code,200)
            self.assertEqual((await self.client.delete(f"/password-login/verification/{job['id']}")).status_code,200)
        context.close.assert_awaited_once()

    async def test_manual_renew_dispatches_without_writing_or_enabling_account(self):
        from common.services.account_renewal import AccountRenewalClient
        async with self.factory() as db:
            account=await db.get(XYAccount,1);account.status='inactive';await db.commit()
        with patch.object(AccountRenewalClient,'renew',AsyncMock(return_value={'status':'skipped','reason':'account_disabled'})) as renew:
            response=await self.client.post('/cookies/renew-login',json=['fixture'])
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['data']['results'][0]['status'],'skipped')
        renew.assert_not_awaited()
        async with self.factory() as db:
            account=await db.get(XYAccount,1)
            self.assertEqual((account.status,account.cookie,account.login_password),('inactive','unb=101; token=old','secret'))

    async def test_manual_renew_returns_executor_result_and_checks_owner(self):
        from common.services.account_renewal import AccountRenewalClient
        with patch.object(AccountRenewalClient,'renew',AsyncMock(return_value={'status':'verified','reason':''})) as renew:
            response=await self.client.post('/cookies/renew-login',json=['fixture'])
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['data']['results'][0]['status'],'verified')
        renew.assert_awaited_once()
        self.owner=8
        with patch.object(AccountRenewalClient,'renew',AsyncMock()) as renew:
            response=await self.client.post('/cookies/renew-login',json=['fixture'])
        renew.assert_not_awaited()
        self.assertFalse(response.json()['data']['results'][0]['success'])

    async def test_browser_poll_and_cancel_forward_service_auth_over_http(self):
        import asyncio
        import json
        import time
        from app.api.routes import password_login
        self.app.include_router(password_login.router)
        received = []
        async def worker(reader, writer):
            try:
                request = await reader.readuntil(b'\r\n\r\n')
                received.append(request.decode())
                body = json.dumps({'status':'not_found','success':True}).encode()
                writer.write(b'HTTP/1.1 200 OK\r\nContent-Type: application/json\r\nConnection: close\r\nContent-Length: '
                             + str(len(body)).encode() + b'\r\n\r\n' + body)
                await writer.drain()
            finally:
                writer.close()
                await writer.wait_closed()
        server = await asyncio.start_server(worker, '127.0.0.1', 0)
        sid = 'fixture-internal-session'
        password_login._BROWSER_OWNERS[sid] = (self.owner, time.time()+60)
        try:
            url = f"http://127.0.0.1:{server.sockets[0].getsockname()[1]}"
            with patch.object(password_login.settings, 'websocket_service_url', url):
                response = await self.client.get(f'/password-login/check/{sid}')
                self.assertEqual(response.status_code, 200)
                response = await self.client.delete(f'/password-login/cancel/{sid}')
                self.assertEqual(response.status_code, 200)
            self.assertEqual(len(received), 2)
            for request in received:
                self.assertIn('x-internal-token: '+ 'fixture-'*8, request.lower())
        finally:
            password_login._BROWSER_OWNERS.pop(sid, None)
            server.close()
            await server.wait_closed()

    async def test_request_policy_can_be_saved_read_and_fences_stale_edit(self):
        before = await self.client.get('/cookies/fixture/request-policy')
        self.assertEqual(before.status_code, 200)
        self.assertFalse(before.json()['data']['configured'])
        saved = await self.client.put('/cookies/fixture/request-policy', json={
            'expected_config_version': 0, 'min_interval_seconds': 10,
            'max_interval_seconds': 15, 'requests_per_minute': 4})
        self.assertEqual(saved.status_code, 200, saved.text)
        data = saved.json()['data']
        self.assertEqual(data['config_version'], 1)
        self.assertEqual(data['effective_interval_seconds'], 15)
        self.assertEqual(data['pending_consumers'], ['websocket', 'scheduler'])
        self.assertNotIn('secret', saved.text)
        read = (await self.client.get('/cookies/fixture/request-policy')).json()['data']
        self.assertEqual(read['values'], {'min_interval_seconds':10, 'max_interval_seconds':15, 'requests_per_minute':4})
        stale = await self.client.put('/cookies/fixture/request-policy', json={
            'expected_config_version': 0, 'min_interval_seconds': 1})
        self.assertEqual(stale.status_code, 409)
        self.owner = 8
        self.assertEqual((await self.client.get('/cookies/fixture/request-policy')).status_code, 404)

    async def test_request_policy_rejects_invalid_intervals_and_keeps_other_config(self):
        from common.services import account_policy
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            state = account_policy.snapshot(account)
            state['config_values'] = {'setting_false': False, 'setting_zero': 0, 'setting_null': None}
            state.update(business_state='verification_required', reason='verification_required')
            account_policy.store(account, state)
            await session.commit()
        for fields in ({'min_interval_seconds':0}, {'min_interval_seconds':True},
                       {'min_interval_seconds':10,'max_interval_seconds':5}):
            response = await self.client.put('/cookies/fixture/request-policy', json={'expected_config_version':0, **fields})
            self.assertEqual(response.status_code, 422)
        response = await self.client.put('/cookies/fixture/request-policy', json={'expected_config_version':0,'min_interval_seconds':10})
        self.assertEqual(response.status_code, 200, response.text)
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            state = account_policy.snapshot(account)
            self.assertEqual(state['config_values']['setting_zero'], 0)
            self.assertIs(state['config_values']['setting_false'], False)
            self.assertIsNone(state['config_values']['setting_null'])
            self.assertEqual(state['business_state'], 'verification_required')
            self.assertEqual(account.login_password, 'secret')

    async def test_configuration_reload_keeps_partial_results_and_retries_only_missing_consumer(self):
        from aiohttp import web
        from common.services.account_configuration import ConfigurationStore
        from app.core.config import get_settings
        from common.utils.internal_auth import is_valid_internal_token
        calls=[]
        failed={'websocket':True}
        async def consumer(request):
            name=request.match_info['consumer']
            self.assertTrue(is_valid_internal_token(request.headers.get('X-Internal-Token'),'fixture-'*8))
            calls.append(name)
            if failed.get(name):
                return web.json_response({'detail':{'code':'offline'}}, status=503)
            data=await request.json()
            result=await ConfigurationStore(self.factory).acknowledge(data['owner_id'], data['account_id'], data['config_version'], name)
            return web.json_response(result)
        app=web.Application();app.router.add_post('/{consumer}/internal/account-configuration',consumer)
        runner=web.AppRunner(app);await runner.setup()
        site=web.TCPSite(runner,'127.0.0.1',0);await site.start()
        base=f'http://127.0.0.1:{site._server.sockets[0].getsockname()[1]}'
        try:
            saved=await self.client.put('/cookies/fixture/request-policy', json={'expected_config_version':0,'min_interval_seconds':10})
            self.assertTrue(saved.json()['success'])
            with patch.object(get_settings(),'websocket_service_url',base+'/websocket'), patch.object(get_settings(),'scheduler_service_url',base+'/scheduler'):
                first=await self.client.post('/cookies/fixture/configuration/reload', json={'config_version':1})
                self.assertEqual(first.status_code,200, first.text)
                self.assertFalse(first.json()['data']['complete'])
                self.assertEqual(first.json()['data']['pending_consumers'],['websocket'])
                failed['websocket']=False
                second=await self.client.post('/cookies/fixture/configuration/reload', json={'config_version':1})
                self.assertTrue(second.json()['data']['complete'],second.text)
                self.assertEqual(calls.count('scheduler'),1)
                self.assertEqual(calls.count('websocket'),2)
                stale=await self.client.post('/cookies/fixture/configuration/reload', json={'config_version':0})
                self.assertEqual(stale.status_code,409)
                self.assertEqual(len(calls),3)
        finally:
            await runner.cleanup()

    async def test_delete_rejects_outstanding_facts_without_stopping_the_account(self):
        from sqlalchemy import insert, delete
        from common.models.delivery_intent import DeliveryIntent
        from common.models.account_operation import AccountOperation
        from common.models.reply_state import reply_outbox
        await self.create_deletion_tables()
        cases = (
            (DeliveryIntent.__table__, {'id':'delivery', 'owner_id':7, 'account_id':'fixture',
             'order_no':'completed-order', 'card_id':1, 'card_type':'text', 'quantity':1,
             'content_state':'confirmed', 'confirm_state':'unknown'}, 'unfinished_deliveries'),
            (reply_outbox, {'id':'reply', 'account_id':'fixture', 'chat_id':'chat', 'request_id':'r',
             'content_hash':'hash', 'status':'unknown', 'result':{}}, 'pending_replies'),
            (AccountOperation.__table__, {'id':'operation', 'request_id':'op', 'owner_id':7,
             'account_id':'fixture', 'command':'platform_request', 'fingerprint':'hash', 'generation':0,
             'credential_version':0, 'config_version':0, 'status':'submitted', 'deadline':0}, 'pending_operations'),
        )
        for table, values, field in cases:
            with self.subTest(field=field):
                async with self.factory() as session:
                    await session.execute(insert(table).values(**values))
                    await session.commit()
                with patch('app.services.websocket_client.websocket_client.stop_account', AsyncMock()) as stop:
                    preview = await self.client.get('/cookies/fixture/delete-preview')
                    self.assertEqual(preview.status_code, 200, preview.text)
                    self.assertFalse(preview.json()['data']['can_delete'])
                    self.assertEqual(preview.json()['data'][field], 1)
                    rejected = await self.client.delete('/cookies/fixture')
                    self.assertEqual(rejected.status_code, 409, rejected.text)
                    stop.assert_not_called()
                async with self.factory() as session:
                    self.assertEqual((await session.get(XYAccount, 1)).status, 'active')
                    await session.execute(delete(table))
                    await session.commit()

    async def test_manual_new_account_creates_queryable_credential_job_before_start(self):
        from sqlalchemy import select
        from common.services.account_policy import snapshot
        await self.create_owner()
        with patch('app.services.account_jobs.async_session_maker', self.factory), \
             patch('app.services.account_jobs.validate_credentials', AsyncMock(return_value='unb=202; token=verified')), \
             patch('app.services.account_jobs.restart_account', AsyncMock()) as restart, \
             patch('app.services.websocket_client.websocket_client.start_account', AsyncMock(return_value={'success':True})) as start:
            response = await self.client.post('/cookies', json={'id':'new-account', 'value':'unb=202; token=raw'})
        self.assertEqual(response.status_code, 202, response.text)
        self.assertNotIn('token=', response.text)
        job = response.json()['data']['job']
        self.assertEqual(job['status'], 'processing')
        query = await self.client.get('/cookies/new-account/credential-jobs/' + job['id'])
        self.assertEqual(query.json()['data']['status'], 'verified')
        start.assert_not_called()
        restart.assert_awaited_once_with('new-account')
        async with self.factory() as session:
            account = await session.scalar(select(XYAccount).where(XYAccount.account_id=='new-account'))
            self.assertEqual(account.unb, '202')
            self.assertEqual(account.cookie, 'unb=202; token=verified')
            self.assertEqual(snapshot(account)['business_state'], 'unchecked')

    async def test_manual_new_alias_for_existing_platform_identity_is_rejected(self):
        await self.create_owner()
        with patch('app.services.websocket_client.websocket_client.start_account', AsyncMock(return_value={'success':True})) as start:
            response = await self.client.post('/cookies', json={'id':'duplicate-alias', 'value':'unb=101; token=other'})
        self.assertEqual(response.status_code, 409, response.text)
        start.assert_not_called()

    async def test_excel_masked_credentials_keep_saved_secrets_and_manual_disable(self):
        from io import BytesIO
        from openpyxl import Workbook
        from common.services.account_policy import snapshot
        await self.create_owner()
        from common.models.card import Card
        async with self.engine.begin() as connection:
            await connection.run_sync(Card.__table__.create)
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            account.status, account.disable_reason = 'disabled', 'manual'
            account.remark, account.proxy_pass = 'keep remark', 'proxy-secret'
            await session.commit()
        book = Workbook(); sheet = book.active; sheet.title='账号基本信息'
        sheet.append(['账号ID','Cookie','状态','用户名','登录密码','代理密码','备注'])
        sheet.append(['fixture','********','active','','********','********',''])
        output = BytesIO(); book.save(output)
        validate = AsyncMock(return_value='unb=101; token=verified')
        with patch('app.services.account_jobs.async_session_maker', self.factory), \
             patch('app.services.account_jobs.validate_credentials', validate), \
             patch('app.services.account_jobs.restart_account', AsyncMock()) as restart, \
             patch('app.services.websocket_client.websocket_client.start_account', AsyncMock(return_value={'success':True})) as start:
            response = await self.client.post('/cookies/import', files={'file':('accounts.xlsx',output.getvalue())}, data={'enable_all':'false'})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()['data']
        self.assertEqual(data['submitted'], 1, response.text)
        self.assertEqual(data['started'], 0)
        self.assertNotIn('secret', response.text)
        validate.assert_awaited_once()
        self.assertEqual(validate.await_args.args[0], 'unb=101; token=old')
        start.assert_not_called(); restart.assert_not_called()
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            self.assertEqual((account.status, account.disable_reason), ('disabled','manual'))
            self.assertEqual((account.username, account.login_password, account.proxy_pass), ('seller','secret','proxy-secret'))
            self.assertEqual(account.remark, 'keep remark')
            self.assertEqual(account.cookie, 'unb=101; token=verified')
            self.assertEqual(snapshot(account)['business_state'], 'disabled')

    async def test_rejected_import_profile_does_not_apply_other_account_sheets(self):
        from io import BytesIO
        from openpyxl import Workbook
        from common.models.card import Card
        from common.models.default_reply import DefaultReply
        await self.create_owner()
        async with self.engine.begin() as connection:
            for table in (Card.__table__, DefaultReply.__table__):
                await connection.run_sync(table.create)
        async with self.factory() as session:
            account = await session.get(XYAccount, 1)
            account.auto_polish = False
            session.add(DefaultReply(id=1, account_id='fixture', enabled=True, reply_content='original'))
            await session.commit()
        book = Workbook(); basic = book.active; basic.title='账号基本信息'
        basic.append(['账号ID', 'Cookie']); basic.append(['fixture', 'unb=999; token=wrong-identity'])
        switches = book.create_sheet('账号开关配置')
        switches.append(['账号ID', '商品擦亮']); switches.append(['fixture', '是'])
        defaults = book.create_sheet('默认回复')
        defaults.append(['账号ID', '启用', '回复内容']); defaults.append(['fixture', '是', 'overwritten'])
        output=BytesIO(); book.save(output)
        response = await self.client.post('/cookies/import', files={'file':('accounts.xlsx',output.getvalue())})
        self.assertEqual(response.json()['data']['submitted'], 0)
        self.assertEqual(response.json()['data']['failed'], 1, response.text)
        async with self.factory() as session:
            self.assertFalse((await session.get(XYAccount, 1)).auto_polish)
            self.assertEqual((await session.get(DefaultReply, 1)).reply_content, 'original')

    async def test_import_validated_ai_settings_preserve_keys_and_rollback_bad_workbook(self):
        from io import BytesIO
        from openpyxl import Workbook
        from common.services.account_policy import snapshot
        import json
        await self.create_owner()
        async with self.factory() as session:
            account = await session.get(XYAccount,1)
            account.remark, account.proxy_user = 'original', 'proxy-login'
            account.metadata_json = {'ai_reply_settings': {'api_key':'saved-model-key', 'ai_enabled':False}}
            await session.commit()
        async def upload(provider):
            book=Workbook(); basic=book.active; basic.title='账号基本信息'
            basic.append(['账号ID','Cookie','备注','代理用户名'])
            basic.append(['fixture','********','new remark','********'])
            ai=book.create_sheet('AI回复设置')
            ai.append(['账号ID','AI回复设置JSON'])
            ai.append(['fixture', json.dumps({'api_key':'********','provider_type':provider})])
            data=BytesIO(); book.save(data); book.close()
            with patch('app.services.account_jobs.async_session_maker',self.factory), \
                 patch('app.services.account_jobs.validate_credentials',AsyncMock(return_value='unb=101; token=verified')), \
                 patch('app.services.account_jobs.restart_account',AsyncMock()):
                return await self.client.post('/cookies/import',files={'file':('accounts.xlsx',data.getvalue())})
        bad=await upload('not-a-supported-protocol')
        self.assertFalse(bad.json()['success'], bad.text)
        self.assertEqual(bad.json()['data']['submitted'],0)
        self.assertEqual(bad.json()['data']['updated'],0)
        async with self.factory() as session:
            account=await session.get(XYAccount,1)
            self.assertEqual(account.remark,'original')
            self.assertEqual(snapshot(account)['config_version'],0)
        good=await upload('openai')
        self.assertTrue(good.json()['success'],good.text)
        async with self.factory() as session:
            account=await session.get(XYAccount,1)
            self.assertEqual(account.proxy_user,'proxy-login')
            self.assertEqual(account.metadata_json['ai_reply_settings']['api_key'],'saved-model-key')
            self.assertEqual(account.metadata_json['ai_reply_settings']['config_version'],1)
            self.assertEqual(account.remark,'new remark')

    async def test_invalid_legacy_frequency_is_visible_without_nonfinite_json(self):
        from common.services.account_policy import snapshot,store
        async with self.factory() as session:
            account=await session.get(XYAccount,1)
            state=snapshot(account);state['config_values']={'risk':{'min_interval_seconds':float('nan')}}
            store(account,state);await session.commit()
        response=await self.client.get('/cookies/fixture/request-policy')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['data']['error'],'risk_invalid')
        self.assertEqual(response.json()['data']['values'],{})

    async def test_excel_upload_size_limit_precedes_workbook_parsing(self):
        response=await self.client.post('/cookies/import',files={'file':('accounts.xlsx',b'x'*(10*1024*1024+1))})
        self.assertFalse(response.json()['success'])
        self.assertEqual(response.json()['message'],'import_file_too_large')
