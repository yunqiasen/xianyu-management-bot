"""S1 Excel真实HTTP导入导出：图片往返、坏行局部失败、不删除既有规则。"""
import io
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock
from sqlalchemy import BigInteger, select
from sqlalchemy.ext.compiler import compiles
from openpyxl import Workbook, load_workbook
import api_suite

@compiles(BigInteger, 'sqlite')
def bigint_sqlite(type_, compiler, **kw): return 'INTEGER'

class KeywordTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        await api_suite.ManagementApiTests.asyncSetUp(self)
        from app.api.routes import keywords
        from app.services.keyword_service import KeywordService
        from common.models.xy_keyword_rule import XYKeywordRule
        from common.models.xy_catalog_item import XYCatalogItem
        self.rules = XYKeywordRule
        async with self.engine.begin() as conn:
            await conn.run_sync(lambda c: self.api.reply_events.metadata.create_all(c, tables=[XYKeywordRule.__table__, XYCatalogItem.__table__]))
        self.session = self.sessions()
        self.account = SimpleNamespace(id=1, owner_id=1, account_id='a')
        self.app.include_router(keywords.router, prefix='/keywords')
        self.app.dependency_overrides[keywords.deps.get_current_active_user] = lambda: SimpleNamespace(id=1, role='user')
        self.app.dependency_overrides[keywords.deps.get_account_service] = lambda: SimpleNamespace(get_account_for_user=AsyncMock(return_value=self.account))
        self.app.dependency_overrides[keywords.deps.get_keyword_service] = lambda: KeywordService(self.session)

    async def asyncTearDown(self):
        await self.session.close()
        await api_suite.ManagementApiTests.asyncTearDown(self)

    async def test_image_excel_roundtrip_reports_bad_rows_and_preserves_rule_id(self):
        wb = Workbook(); ws=wb.active
        ws.append(['关键词','商品ID','关键词内容','回复类型','图片地址'])
        ws.append(['图','item','图片说明','image','https://fixture.example/image.png'])
        ws.append(['坏','item','no','unknown',''])
        ws.append(['空回复','','','text',''])
        output=io.BytesIO(); wb.save(output)
        response=await self.client.post('/keywords/a/import', files={'file':('rules.xlsx', output.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        data=response.json()['data']
        self.assertEqual(data['added'],2)
        self.assertEqual(data['errors'][0]['row'],3)
        rows=(await self.session.execute(select(self.rules))).scalars().all()
        self.assertEqual(len(rows),2)
        ids={row.keyword:row.id for row in rows}
        response=await self.client.get('/keywords/a/export')
        self.assertEqual(response.status_code,200)
        exported=list(load_workbook(io.BytesIO(response.content)).active.values)
        self.assertIn('图片地址',exported[0])
        self.assertEqual(len(exported),3)
        response=await self.client.post('/keywords/a/import', files={'file':('rules.xlsx',response.content)})
        self.assertEqual(response.json()['data']['updated'],2)
        self.session.expire_all()
        self.assertEqual({row.keyword:row.id for row in (await self.session.execute(select(self.rules))).scalars()},ids)

if __name__ == '__main__': unittest.main()
