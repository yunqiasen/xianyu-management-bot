"""
账号数据导入服务

功能：
1. 解析导出的Excel文件（openpyxl）
2. 按Sheet逐一导入数据到数据库（upsert逻辑）
3. 支持两种模式：保存（按Excel状态）/ 保存并全部启用
4. 返回导入统计结果
"""
from __future__ import annotations

import json
from io import BytesIO
from copy import deepcopy
from zipfile import ZipFile

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 50000
from typing import Any

from loguru import logger
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from common.models.xy_account import XYAccount
from common.models.card import Card
from common.models.card_item_relation import CardItemRelation
from common.models.xy_keyword_rule import XYKeywordRule
from common.models.default_reply import DefaultReply
from common.models.xy_catalog_item import XYCatalogItem
from common.models.confirm_receipt_message import ConfirmReceiptMessage
from common.models.auto_rate_config import AutoRateConfig
from common.utils.time_utils import get_beijing_now_naive


def _parse_bool(value: str | None) -> bool:
    """解析布尔值：'是'/True/'true'/'1' → True，其余 → False"""
    if value is None:
        return False
    v = str(value).strip().lower()
    return v in ("是", "true", "1", "yes")


def _parse_int(value: str | None, default: int = 0) -> int:
    """解析整数，失败返回默认值"""
    if value is None or str(value).strip() == "":
        return default
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return default


def _parse_str(value: Any) -> str:
    """解析字符串，None → 空字符串"""
    if value is None:
        return ""
    return str(value).strip()


def _parse_json(value: str | None) -> Any:
    """解析JSON字符串，失败返回None"""
    if not value or str(value).strip() == "":
        return None
    try:
        return json.loads(str(value))
    except (json.JSONDecodeError, TypeError):
        return None


def _read_sheet_rows(wb, sheet_name: str) -> list[dict[str, str]]:
    """读取Sheet为字典列表（表头作为key）"""
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    if (ws.max_column or 0) > 128 or (ws.max_row or 0) > MAX_IMPORT_ROWS:
        raise ValueError('import_sheet_too_large')
    rows = []
    for row in ws.iter_rows(values_only=True):
        if len(rows) >= MAX_IMPORT_ROWS or len(row)>128:
            raise ValueError('import_sheet_too_large')
        rows.append(row)
    if len(rows) < 2:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    result = []
    for row in rows[1:]:
        row_dict = {}
        for i, header in enumerate(headers):
            if not header:
                continue
            val = row[i] if i < len(row) else None
            row_dict[header] = str(val).strip() if val is not None else ""
        # 跳过全空行
        if any(v for v in row_dict.values()):
            result.append(row_dict)
    return result


class AccountImportService:
    """账号数据导入服务"""

    def __init__(self, session: AsyncSession, owner_id: int):
        self.session = session
        self.owner_id = owner_id
        self.inserted = 0
        self.updated = 0
        self.started = 0
        self.failed = 0
        self.errors: list[str] = []
        self.pending_credential_jobs: list[tuple[str, str, str]] = []
        # 导入过程中的映射缓存
        self._account_id_to_pk: dict[str, int] = {}
        self._card_name_spec_to_id: dict[str, int] = {}

    async def import_accounts(
        self,
        file_content: bytes,
        enable_all: bool = False,
    ) -> dict:
        """导入账号数据

        Args:
            file_content: Excel文件内容
            enable_all: 是否全部启用

        Returns:
            导入结果统计
        """
        if len(file_content) > MAX_IMPORT_BYTES:
            return {'success':False,'message':'import_file_too_large','data':None}
        try:
            with ZipFile(BytesIO(file_content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 64*1024*1024:
                    return {'success':False,'message':'import_workbook_too_large','data':None}
            wb = load_workbook(BytesIO(file_content), read_only=True, data_only=True)
            if len(wb.sheetnames)>24:
                wb.close()
                return {'success':False,'message':'import_too_many_sheets','data':None}
        except Exception as e:
            return {
                "success": False,
                "message": "Excel文件格式错误",
                "data": None,
            }

        try:
            # 按顺序导入各Sheet
            await self._import_accounts_basic(wb, enable_all)
            await self._import_account_switches(wb)
            await self._import_ai_settings(wb)
            await self._import_catalog_items(wb)
            await self._import_cards(wb)
            await self._import_card_item_relations(wb)
            await self._import_keyword_rules(wb)
            await self._import_default_replies(wb)
            await self._import_message_filters(wb)
            await self._import_confirm_receipt(wb)
            await self._import_auto_rate(wb)
            await self.session.commit()

        except Exception as e:
            await self.session.rollback()
            logger.error("导入过程异常: {}", type(e).__name__)
            self.errors.append("部分配置保存失败，请检查导入文件")
            self.failed += 1
            self.pending_credential_jobs.clear()
            self.inserted = self.updated = 0
        finally:
            wb.close()

        message = (
            f"导入完成：新增 {self.inserted} 个，更新 {self.updated} 个，"
            f"凭据检查提交 {len(self.pending_credential_jobs)} 个，失败 {self.failed} 个"
        )
        return {
            "success": not self.errors,
            "partial": bool(self.errors) and bool(self.inserted or self.updated),
            "message": message,
            "data": {
                "inserted": self.inserted,
                "updated": self.updated,
                "started": 0,
                "submitted": len(self.pending_credential_jobs),
                "credential_jobs": [{'account_id':account_id, 'id':job_id, 'status':'processing'}
                                    for account_id, job_id, _ in self.pending_credential_jobs],
                "failed": self.failed,
                "errors": self.errors[:20],  # 最多返回20条错误
            },
        }

    # ==================== 各Sheet导入逻辑 ====================

    async def _import_accounts_basic(self, wb, enable_all: bool) -> None:
        from app.services.account_service import AccountService
        service = AccountService(self.session)
        seen = set()
        for row in _read_sheet_rows(wb, "账号基本信息"):
            account_id = _parse_str(row.get('账号ID'))
            if not account_id:
                continue
            if account_id in seen:
                self.failed += 1
                self.errors.append(f'账号 {account_id}: 同一文件存在重复行')
                continue
            seen.add(account_id)
            try:
                async with self.session.begin_nested():
                    profile = {}
                    string_fields = {'备注':'remark', '用户名':'username', '登录密码':'login_password',
                        '代理类型':'proxy_type', '代理地址':'proxy_host', '代理用户名':'proxy_user', '代理密码':'proxy_pass'}
                    for source, target in string_fields.items():
                        value = _parse_str(row.get(source))
                        if value:
                            profile[target] = value
                    for source, target in {'暂停时长(秒)':'pause_duration', '相同消息等待时间(秒)':'message_expire_time', '代理端口':'proxy_port'}.items():
                        if _parse_str(row.get(source)):
                            value = int(str(row[source]).strip())
                            if value < 0:
                                raise ValueError('时长或端口格式错误')
                            profile[target] = value
                    if _parse_str(row.get('显示浏览器')):
                        profile['show_browser'] = _parse_bool(row['显示浏览器'])
                    prior = await service.get_account_for_user(self.owner_id, account_id)
                    source_status = _parse_str(row.get('状态')).lower()
                    enabled = True if enable_all else (False if source_status in {'inactive','disabled','suspended'} else None)
                    if prior is None and enabled is None:
                        enabled = True
                    account, job, candidate, created = await service.stage_cookie_import(
                        self.owner_id, account_id, row.get('Cookie'), enabled=enabled, profile=profile, login_method='import', commit=False)
                    self._account_id_to_pk[account_id] = account.id
                    self.pending_credential_jobs.append((account_id, job['id'], candidate))
                    self.inserted += int(created)
                    self.updated += int(not created)
            except ValueError as exc:
                self.failed += 1
                self.errors.append(f'账号 {account_id}: {exc}')
            except Exception as exc:
                self.failed += 1
                self.errors.append(f'账号 {account_id}: 资料保存失败')
                logger.warning('账号导入行失败: {}', type(exc).__name__)

    async def _import_account_switches(self, wb) -> None:
        """导入账号开关配置"""
        rows = _read_sheet_rows(wb, "账号开关配置")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            if account_id not in self._account_id_to_pk:
                continue
            stmt = select(XYAccount).where(
                XYAccount.account_id == account_id,
                XYAccount.owner_id == self.owner_id,
            )
            result = await self.session.execute(stmt)
            account = result.scalars().first()
            if not account:
                continue

            fields = {
                "自动确认收货":"auto_confirm", "定时补发货":"scheduled_redelivery",
                "定时补评价":"scheduled_rate", "商品擦亮":"auto_polish",
                "发货成功再发卡券":"confirm_before_send", "卡券发送成功再确认发货":"send_before_confirm",
                "只发卡券不确认发货":"only_send_card", "自动求小红花":"auto_red_flower",
                "禁止发货":"delivery_disabled", "主动关闭订单":"auto_close_order",
                "关闭后发卡券":"delivery_only_card_after_close",
            }
            explicit = []
            for label, field in fields.items():
                if label in row:
                    setattr(account, field, _parse_bool(row[label]))
                    explicit.append(field)
            if account.only_send_card:
                account.auto_confirm = False
                account.confirm_before_send = False
                account.send_before_confirm = False
            elif account.confirm_before_send:
                account.send_before_confirm = False
            if "禁止发货原因" in row:
                account.delivery_disabled_reason = _parse_str(row["禁止发货原因"]) or None
                explicit.append("delivery_disabled_reason")
            if "禁止发货排除商品" in row:
                excluded = _parse_json(row["禁止发货排除商品"])
                if isinstance(excluded, list):
                    account.delivery_disabled_excluded_items = excluded
            if any(label in row for label in ('禁止发货', '禁止发货原因', '禁止发货排除商品', '主动关闭订单', '关闭后发卡券')):
                from common.services.delivery_rule_configuration import sync_legacy_credit_rule
                explicit.extend(await sync_legacy_credit_rule(self.session, account))
            # All sheets are one transaction/configuration batch, already fenced
            # by the profile stage. Do not invalidate its pending credential job.
            from common.services.typed_settings import clear_inheritance
            clear_inheritance(account, explicit)
            self.session.add(account)

        await self.session.flush()

    async def _import_ai_settings(self, wb) -> None:
        """导入AI回复设置"""
        rows = _read_sheet_rows(wb, "AI回复设置")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            if account_id not in self._account_id_to_pk:
                continue
            ai_json = _parse_json(row.get("AI回复设置JSON"))
            if not ai_json:
                continue
            stmt = select(XYAccount).where(
                XYAccount.account_id == account_id,
                XYAccount.owner_id == self.owner_id,
            )
            result = await self.session.execute(stmt)
            account = result.scalars().first()
            if not account:
                continue

            from app.services.ai_reply_service import AIReplySettingsService, merge_ai_settings
            if not isinstance(ai_json, dict):
                raise ValueError('AI配置必须为对象')
            existing = AIReplySettingsService(self.session)._extract_settings(account)
            merged = merge_ai_settings(existing, ai_json)
            merged['config_version'] = int(existing.get('config_version') or 0) + 1
            metadata = deepcopy(account.metadata_json or {})
            metadata["ai_reply_settings"] = merged
            account.metadata_json = metadata
            self.session.add(account)

        await self.session.flush()

    async def _import_catalog_items(self, wb) -> None:
        """导入商品目录"""
        rows = _read_sheet_rows(wb, "商品目录")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            item_id = _parse_str(row.get("商品ID"))
            if not account_id or not item_id:
                continue
            account_pk = self._account_id_to_pk.get(account_id)
            if not account_pk:
                continue

            stmt = select(XYCatalogItem).where(
                XYCatalogItem.account_pk == account_pk,
                XYCatalogItem.item_id == item_id,
            )
            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.title = _parse_str(row.get("标题")) or existing.title
                existing.price = _parse_str(row.get("价格")) or existing.price
                existing.ai_prompt = _parse_str(row.get("AI提示词")) or existing.ai_prompt
            else:
                item = XYCatalogItem(
                    owner_id=self.owner_id,
                    account_pk=account_pk,
                    item_id=item_id,
                    title=_parse_str(row.get("标题")),
                    price=_parse_str(row.get("价格")),
                    ai_prompt=_parse_str(row.get("AI提示词")) or None,
                    created_at=get_beijing_now_naive(),
                )
                self.session.add(item)

        await self.session.flush()

    async def _import_cards(self, wb) -> None:
        """导入卡券"""
        rows = _read_sheet_rows(wb, "卡券")
        for row in rows:
            name = _parse_str(row.get("卡券名称"))
            if not name:
                continue
            card_type = _parse_str(row.get("类型")) or "text"
            spec_value = _parse_str(row.get("规格值"))

            # 按 名称+规格值 查找
            stmt = select(Card).where(
                Card.user_id == self.owner_id,
                Card.name == name,
            )
            if spec_value:
                stmt = stmt.where(Card.spec_value == spec_value)
            else:
                stmt = stmt.where((Card.spec_value.is_(None)) | (Card.spec_value == ""))

            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.type = card_type
                existing.enabled = _parse_bool(row.get("启用"))
                existing.delay_seconds = _parse_int(row.get("延迟秒数"), 0)
                existing.is_multi_spec = _parse_bool(row.get("多规格"))
                existing.spec_name = _parse_str(row.get("规格名")) or existing.spec_name
                existing.spec_value = spec_value or existing.spec_value
                existing.api_config = _parse_str(row.get("API配置")) or existing.api_config
                existing.text_content = _parse_str(row.get("文本内容")) or existing.text_content
                existing.data_content = _parse_str(row.get("数据内容")) or existing.data_content
                existing.image_url = _parse_str(row.get("图片URL")) or existing.image_url
                existing.image_urls = _parse_str(row.get("多图片URL")) or existing.image_urls
                self.session.add(existing)
                self._card_name_spec_to_id[f"{name}|{spec_value}"] = existing.id
            else:
                card = Card(
                    user_id=self.owner_id,
                    name=name,
                    type=card_type,
                    enabled=_parse_bool(row.get("启用")),
                    delay_seconds=_parse_int(row.get("延迟秒数"), 0),
                    is_multi_spec=_parse_bool(row.get("多规格")),
                    spec_name=_parse_str(row.get("规格名")) or None,
                    spec_value=spec_value or None,
                    api_config=_parse_str(row.get("API配置")) or None,
                    text_content=_parse_str(row.get("文本内容")) or None,
                    data_content=_parse_str(row.get("数据内容")) or None,
                    image_url=_parse_str(row.get("图片URL")) or None,
                    image_urls=_parse_str(row.get("多图片URL")) or None,
                )
                self.session.add(card)
                await self.session.flush()
                self._card_name_spec_to_id[f"{name}|{spec_value}"] = card.id

        await self.session.flush()

        if "卡券商品关联" not in wb.sheetnames:
            return
        # 补充映射：查询所有该用户的卡券
        stmt = select(Card.id, Card.name, Card.spec_value).where(Card.user_id == self.owner_id)
        result = await self.session.execute(stmt)
        for card_id, card_name, card_spec in result.all():
            key = f"{card_name}|{card_spec or ''}"
            self._card_name_spec_to_id[key] = card_id

    async def _import_card_item_relations(self, wb) -> None:
        """导入卡券商品关联"""
        rows = _read_sheet_rows(wb, "卡券商品关联")
        for row in rows:
            card_name = _parse_str(row.get("卡券名称"))
            item_id = _parse_str(row.get("商品ID"))
            if not card_name or not item_id:
                continue

            # 查找卡券ID（先精确匹配名称，不带规格）
            card_id = None
            for key, cid in self._card_name_spec_to_id.items():
                if key.startswith(f"{card_name}|"):
                    card_id = cid
                    break
            if not card_id:
                continue

            # 检查是否已存在
            stmt = select(CardItemRelation).where(
                CardItemRelation.card_id == card_id,
                CardItemRelation.item_id == item_id,
                CardItemRelation.user_id == self.owner_id,
            )
            result = await self.session.execute(stmt)
            if result.scalars().first():
                continue  # 已存在，跳过

            source = _parse_str(row.get("来源")) or "own"
            rel = CardItemRelation(
                user_id=self.owner_id,
                card_id=card_id,
                item_id=item_id,
                source=source,
            )
            self.session.add(rel)

        await self.session.flush()

    async def _import_keyword_rules(self, wb) -> None:
        """导入关键词规则"""
        rows = _read_sheet_rows(wb, "关键词规则")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            keyword = _parse_str(row.get("关键词"))
            if not account_id or not keyword:
                continue
            account_pk = self._account_id_to_pk.get(account_id)
            if not account_pk:
                continue

            item_id = _parse_str(row.get("商品ID")) or None

            stmt = select(XYKeywordRule).where(
                XYKeywordRule.account_pk == account_pk,
                XYKeywordRule.keyword == keyword,
            )
            if item_id:
                stmt = stmt.where(XYKeywordRule.item_id == item_id)
            else:
                stmt = stmt.where((XYKeywordRule.item_id.is_(None)) | (XYKeywordRule.item_id == ""))

            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.reply_content = _parse_str(row.get("回复内容")) or existing.reply_content
                existing.reply_type = _parse_str(row.get("回复类型")) or existing.reply_type
                existing.image_url = _parse_str(row.get("图片URL")) or existing.image_url
                existing.location_name = _parse_str(row.get("定位名称")) or existing.location_name
                existing.location_longitude = _parse_str(row.get("经度")) or existing.location_longitude
                existing.location_latitude = _parse_str(row.get("纬度")) or existing.location_latitude
                existing.location_title = _parse_str(row.get("位置标题")) or existing.location_title
                existing.location_subtitle = _parse_str(row.get("位置副标题")) or existing.location_subtitle
                existing.priority = _parse_int(row.get("优先级"), existing.priority)
                existing.is_active = _parse_bool(row.get("启用"))
            else:
                rule = XYKeywordRule(
                    owner_id=self.owner_id,
                    account_pk=account_pk,
                    keyword=keyword,
                    reply_content=_parse_str(row.get("回复内容")) or None,
                    reply_type=_parse_str(row.get("回复类型")) or "text",
                    image_url=_parse_str(row.get("图片URL")) or None,
                    location_name=_parse_str(row.get("定位名称")) or None,
                    location_longitude=_parse_str(row.get("经度")) or None,
                    location_latitude=_parse_str(row.get("纬度")) or None,
                    location_title=_parse_str(row.get("位置标题")) or None,
                    location_subtitle=_parse_str(row.get("位置副标题")) or None,
                    item_id=item_id,
                    priority=_parse_int(row.get("优先级"), 100),
                    is_active=_parse_bool(row.get("启用")),
                )
                self.session.add(rule)

        await self.session.flush()

    async def _import_default_replies(self, wb) -> None:
        """导入默认回复"""
        rows = _read_sheet_rows(wb, "默认回复")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            if account_id not in self._account_id_to_pk:
                continue
            item_id = _parse_str(row.get("商品ID")) or None

            stmt = select(DefaultReply).where(DefaultReply.account_id == account_id)
            if item_id:
                stmt = stmt.where(DefaultReply.item_id == item_id)
            else:
                stmt = stmt.where((DefaultReply.item_id.is_(None)) | (DefaultReply.item_id == ""))

            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.enabled = _parse_bool(row.get("启用"))
                existing.reply_content = _parse_str(row.get("回复内容")) or existing.reply_content
                existing.reply_image = _parse_str(row.get("回复图片")) or existing.reply_image
                existing.reply_once = _parse_bool(row.get("仅回复一次"))
                existing.reply_type = _parse_str(row.get("回复类型")) or existing.reply_type or "text"
                existing.api_url = _parse_str(row.get("API地址")) or existing.api_url
                api_timeout = _parse_int(row.get("API超时"))
                if api_timeout:
                    existing.api_timeout = api_timeout
                existing.location_name = _parse_str(row.get("\u5b9a\u4f4d\u540d\u79f0")) or existing.location_name
                existing.location_longitude = _parse_str(row.get("\u7ecf\u5ea6")) or existing.location_longitude
                existing.location_latitude = _parse_str(row.get("\u7eac\u5ea6")) or existing.location_latitude
                existing.location_title = _parse_str(row.get("\u4f4d\u7f6e\u6807\u9898")) or existing.location_title
                existing.location_subtitle = _parse_str(row.get("\u4f4d\u7f6e\u526f\u6807\u9898")) or existing.location_subtitle
            else:
                reply = DefaultReply(
                    account_id=account_id,
                    item_id=item_id,
                    enabled=_parse_bool(row.get("启用")),
                    reply_content=_parse_str(row.get("回复内容")) or None,
                    reply_image=_parse_str(row.get("回复图片")) or None,
                    reply_once=_parse_bool(row.get("仅回复一次")),
                    reply_type=_parse_str(row.get("回复类型")) or "text",
                    api_url=_parse_str(row.get("API地址")) or None,
                    api_timeout=_parse_int(row.get("API超时")) or 80,
                )
                reply.location_name = _parse_str(row.get("\u5b9a\u4f4d\u540d\u79f0")) or None
                reply.location_longitude = _parse_str(row.get("\u7ecf\u5ea6")) or None
                reply.location_latitude = _parse_str(row.get("\u7eac\u5ea6")) or None
                reply.location_title = _parse_str(row.get("\u4f4d\u7f6e\u6807\u9898")) or None
                reply.location_subtitle = _parse_str(row.get("\u4f4d\u7f6e\u526f\u6807\u9898")) or None
                self.session.add(reply)

        await self.session.flush()

    async def _import_message_filters(self, wb) -> None:
        """导入消息过滤规则"""
        rows = _read_sheet_rows(wb, "消息过滤规则")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            keyword = _parse_str(row.get("关键词"))
            filter_type = _parse_str(row.get("过滤类型"))
            if account_id not in self._account_id_to_pk or not keyword or not filter_type:
                continue

            # 检查唯一约束
            check_sql = text("""
                SELECT id FROM xy_message_filters
                WHERE account_id = :account_id AND keyword = :keyword AND filter_type = :filter_type
                LIMIT 1
            """)
            result = await self.session.execute(check_sql, {
                "account_id": account_id, "keyword": keyword, "filter_type": filter_type
            })
            if result.scalar_one_or_none():
                continue  # 已存在，跳过

            enabled = _parse_bool(row.get("启用"))
            insert_sql = text("""
                INSERT INTO xy_message_filters (account_id, keyword, filter_type, enabled)
                VALUES (:account_id, :keyword, :filter_type, :enabled)
            """)
            await self.session.execute(insert_sql, {
                "account_id": account_id, "keyword": keyword,
                "filter_type": filter_type, "enabled": 1 if enabled else 0,
            })

        await self.session.flush()

    async def _import_confirm_receipt(self, wb) -> None:
        """导入确认收货消息"""
        rows = _read_sheet_rows(wb, "确认收货消息")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            if account_id not in self._account_id_to_pk:
                continue

            stmt = select(ConfirmReceiptMessage).where(ConfirmReceiptMessage.account_id == account_id)
            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.enabled = _parse_bool(row.get("启用"))
                existing.message_content = _parse_str(row.get("消息内容")) or existing.message_content
                existing.message_image = _parse_str(row.get("消息图片")) or existing.message_image
            else:
                msg = ConfirmReceiptMessage(
                    account_id=account_id,
                    enabled=_parse_bool(row.get("启用")),
                    message_content=_parse_str(row.get("消息内容")) or None,
                    message_image=_parse_str(row.get("消息图片")) or None,
                )
                self.session.add(msg)

        await self.session.flush()

    async def _import_auto_rate(self, wb) -> None:
        """导入自动评价配置"""
        rows = _read_sheet_rows(wb, "自动评价配置")
        for row in rows:
            account_id = _parse_str(row.get("账号ID"))
            if account_id not in self._account_id_to_pk:
                continue

            stmt = select(AutoRateConfig).where(AutoRateConfig.account_id == account_id)
            result = await self.session.execute(stmt)
            existing = result.scalars().first()

            if existing:
                existing.enabled = _parse_bool(row.get("启用"))
                existing.rate_type = _parse_str(row.get("评价类型")) or existing.rate_type
                existing.text_content = _parse_str(row.get("评价内容")) or existing.text_content
                existing.api_url = _parse_str(row.get("API地址")) or existing.api_url
            else:
                cfg = AutoRateConfig(
                    account_id=account_id,
                    enabled=_parse_bool(row.get("启用")),
                    rate_type=_parse_str(row.get("评价类型")) or "text",
                    text_content=_parse_str(row.get("评价内容")) or None,
                    api_url=_parse_str(row.get("API地址")) or None,
                )
                self.session.add(cfg)

        await self.session.flush()

    # ==================== 启动逻辑 ====================

    async def _get_accounts_to_start(self, enable_all: bool) -> list[tuple[str, str]]:
        """获取需要启动的账号列表 [(account_id, cookie)]"""
        stmt = select(XYAccount.account_id, XYAccount.cookie, XYAccount.status).where(
            XYAccount.owner_id == self.owner_id,
            XYAccount.account_id.in_(list(self._account_id_to_pk.keys())),
        )
        result = await self.session.execute(stmt)
        accounts_to_start = []
        for account_id, cookie, status in result.all():
            if not cookie or not cookie.strip():
                continue
            if enable_all or status == "active":
                accounts_to_start.append((account_id, cookie))
        return accounts_to_start

    async def _start_account(self, account_id: str, cookie: str) -> None:
        """启动单个账号的WebSocket任务"""
        from app.services.websocket_client import websocket_client
        result = await websocket_client.start_account(account_id, cookie, self.owner_id)
        if isinstance(result, dict) and not result.get("success", True):
            raise Exception(result.get("message", "启动失败"))
